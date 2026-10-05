import os
import json
import csv
import io
import math
import hashlib
from typing import List, Union, Tuple, Dict, Any, Optional
from langchain_core.documents import Document
from langchain_core.embeddings import Embeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.vectorstores import FAISS
from langchain_ollama import OllamaEmbeddings


class LightweightFallbackEmbeddings(Embeddings):
    """
    Universal, zero-dependency embedding generator.
    Produces 384-dimensional normalized dense vectors using hash-token n-gram projections.
    Guarantees that vector search works 100% offline without requiring Ollama or torch/sentence-transformers.
    """
    def __init__(self, dim: int = 384):
        self.dim = dim

    def _embed(self, text: str) -> List[float]:
        import re
        tokens = re.findall(r"\w+", text.lower())
        vec = [0.0] * self.dim
        if not tokens:
            return [0.001] * self.dim

        for i, t in enumerate(tokens):
            h = int(hashlib.sha256(t.encode()).hexdigest(), 16)
            idx = h % self.dim
            weight = 1.0 + (1.0 / (1.0 + math.log(i + 1)))
            vec[idx] += weight

            # Bigram hash for phrase context
            if i > 0:
                bi = f"{tokens[i-1]}_{t}"
                h_bi = int(hashlib.md5(bi.encode()).hexdigest(), 16)
                idx_bi = h_bi % self.dim
                vec[idx_bi] += 1.5

        norm = math.sqrt(sum(x * x for x in vec)) or 1.0
        return [round(x / norm, 6) for x in vec]

    def embed_documents(self, texts: List[str]) -> List[List[float]]:
        return [self._embed(t) for t in texts]

    def embed_query(self, text: str) -> List[float]:
        return self._embed(text)


def load_file_content(file_path: str) -> List[Document]:
    """
    Universal document loader supporting PDF, DOCX, TXT, MD, CSV, XLSX, JSON, and code files.
    """
    ext = os.path.splitext(file_path)[1].lower()
    filename = os.path.basename(file_path)
    docs = []

    # 1. PDF Loading
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(file_path)
            for page_idx, page in enumerate(reader.pages):
                text = page.extract_text()
                if text and text.strip():
                    docs.append(Document(
                        page_content=text.strip(),
                        metadata={
                            "source_name": filename,
                            "source": file_path,
                            "page": page_idx + 1,
                            "file_type": "PDF",
                        }
                    ))
        except Exception:
            try:
                from langchain_community.document_loaders import PyPDFLoader
                loader = PyPDFLoader(file_path)
                loaded_docs = loader.load()
                for d in loaded_docs:
                    d.metadata["source_name"] = filename
                    d.metadata["file_type"] = "PDF"
                    if "page" in d.metadata:
                        d.metadata["page"] = d.metadata["page"] + 1
                docs.extend(loaded_docs)
            except Exception as e2:
                raise ValueError(f"Failed to load PDF '{filename}': {str(e2)}")

    # 2. DOCX Loading
    elif ext in [".docx", ".doc"]:
        text = ""
        try:
            import docx
            doc = docx.Document(file_path)
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text)
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        full_text.append(f"| {row_text} |")
            text = "\n\n".join(full_text)
        except Exception:
            try:
                import docx2txt
                text = docx2txt.process(file_path)
            except Exception:
                try:
                    import zipfile
                    import xml.etree.ElementTree as ET
                    with zipfile.ZipFile(file_path) as z:
                        xml_content = z.read("word/document.xml")
                        tree = ET.fromstring(xml_content)
                        text_list = [node.text for node in tree.iter() if node.text]
                        text = " ".join(text_list)
                except Exception as e3:
                    raise ValueError(f"Could not read Word document '{filename}': {str(e3)}")

        if text.strip():
            docs.append(Document(
                page_content=text.strip(),
                metadata={
                    "source_name": filename,
                    "source": file_path,
                    "page": 1,
                    "file_type": "Word Document",
                }
            ))

    # 3. CSV / TSV / Excel Loading
    elif ext in [".csv", ".tsv", ".xlsx", ".xls"]:
        try:
            import pandas as pd
            if ext == ".csv":
                df = pd.read_csv(file_path)
            elif ext == ".tsv":
                df = pd.read_csv(file_path, sep="\t")
            else:
                df = pd.read_excel(file_path)
            
            num_rows = len(df)
            columns_str = ", ".join(df.columns.astype(str))
            
            summary_text = (
                f"Dataset: {filename}\n"
                f"Total Rows: {num_rows}\n"
                f"Columns: {columns_str}\n\n"
                f"Data Preview:\n{df.head(10).to_markdown(index=False)}"
            )
            docs.append(Document(
                page_content=summary_text,
                metadata={
                    "source_name": filename,
                    "source": file_path,
                    "page": "Summary",
                    "file_type": "Structured Data",
                }
            ))
            
            chunk_size_rows = 20
            for i in range(0, len(df), chunk_size_rows):
                chunk_df = df.iloc[i : i + chunk_size_rows]
                docs.append(Document(
                    page_content=f"Dataset rows {i+1} to {min(i+chunk_size_rows, len(df))}:\n" + chunk_df.to_markdown(index=False),
                    metadata={
                        "source_name": filename,
                        "source": file_path,
                        "page": f"Rows {i+1}-{min(i+chunk_size_rows, len(df))}",
                        "file_type": "Structured Data",
                    }
                ))
        except Exception:
            try:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    content = f.read()
                    docs.append(Document(
                        page_content=content,
                        metadata={"source_name": filename, "source": file_path, "page": 1, "file_type": "CSV/Data"}
                    ))
            except Exception as e4:
                raise ValueError(f"Failed to load data file '{filename}': {str(e4)}")

    # 4. JSON Loading
    elif ext == ".json":
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                data = json.load(f)
                formatted_json = json.dumps(data, indent=2)
                docs.append(Document(
                    page_content=formatted_json,
                    metadata={
                        "source_name": filename,
                        "source": file_path,
                        "page": 1,
                        "file_type": "JSON Data",
                    }
                ))
        except Exception as e:
            raise ValueError(f"Failed to load JSON file '{filename}': {str(e)}")

    # 5. Plain Text, Markdown, HTML, Code Files
    else:
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
                if content.strip():
                    docs.append(Document(
                        page_content=content.strip(),
                        metadata={
                            "source_name": filename,
                            "source": file_path,
                            "page": 1,
                            "file_type": ext.upper().replace(".", "") or "Text",
                        }
                    ))
        except Exception as e:
            raise ValueError(f"Failed to load file '{filename}': {str(e)}")

    return docs


class SimpleBM25:
    """
    Fast, dependency-free BM25 implementation for hybrid search.
    """
    def __init__(self, documents: List[Document]):
        import re
        self.documents = documents
        self.k1 = 1.5
        self.b = 0.75
        self.doc_len = []
        self.doc_freqs = []
        self.idf = {}
        self.avgdl = 0

        total_len = 0
        df = {}
        for doc in documents:
            tokens = re.findall(r"\w+", doc.page_content.lower())
            self.doc_len.append(len(tokens))
            total_len += len(tokens)
            frequencies = {}
            for t in tokens:
                frequencies[t] = frequencies.get(t, 0) + 1
            self.doc_freqs.append(frequencies)
            for t in frequencies:
                df[t] = df.get(t, 0) + 1

        num_docs = len(documents)
        self.avgdl = (total_len / num_docs) if num_docs > 0 else 0

        for t, freq in df.items():
            self.idf[t] = math.log(1 + (num_docs - freq + 0.5) / (freq + 0.5))

    def search(self, query: str, top_k: int = 4) -> List[Tuple[Document, float]]:
        import re
        tokens = re.findall(r"\w+", query.lower())
        scores = []
        for idx, freqs in enumerate(self.doc_freqs):
            score = 0.0
            doc_l = self.doc_len[idx]
            for t in tokens:
                if t in freqs:
                    tf = freqs[t]
                    idf = self.idf.get(t, 0.1)
                    denom = tf + self.k1 * (1 - self.b + self.b * (doc_l / (self.avgdl or 1)))
                    score += idf * (tf * (self.k1 + 1)) / denom
            scores.append((self.documents[idx], score))

        scores.sort(key=lambda x: x[1], reverse=True)
        return scores[:top_k]


def process_files_to_knowledge_base(
    file_paths: Union[str, List[str]],
    chunk_size: int = 900,
    chunk_overlap: int = 150,
    embedding_model: str = "nomic-embed-text",
) -> Tuple[FAISS, SimpleBM25, List[Document], int, Dict[str, Any]]:
    """
    Universal ingestion pipeline with multi-tiered embedding fallback:
    1. Local Ollama Embeddings (if Ollama is running)
    2. HuggingFace Embeddings (if sentence-transformers installed)
    3. Lightweight Zero-Dependency Fallback Embeddings (always works, 0 dependencies)
    """
    if isinstance(file_paths, str):
        file_paths = [file_paths]

    all_raw_docs: List[Document] = []
    file_stats = {}

    for path in file_paths:
        if not os.path.exists(path):
            continue
        docs = load_file_content(path)
        filename = os.path.basename(path)
        total_chars = sum(len(d.page_content) for d in docs)
        file_stats[filename] = {
            "path": path,
            "pages": len(docs),
            "chars": total_chars,
            "words": sum(len(d.page_content.split()) for d in docs),
            "type": docs[0].metadata.get("file_type", "Document") if docs else "Unknown",
        }
        all_raw_docs.extend(docs)

    if not all_raw_docs:
        raise ValueError("No readable text could be extracted from the uploaded files.")

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", "### ", "## ", "# ", ". ", " ", ""],
        length_function=len,
    )
    chunks = splitter.split_documents(all_raw_docs)

    for idx, c in enumerate(chunks):
        c.metadata["chunk_id"] = idx + 1

    # Tiered Robust Embedding Strategy (Guaranteed never to crash)
    vector_db = None
    try:
        # Tier 1: Try Local Ollama
        embeddings = OllamaEmbeddings(model=embedding_model)
        vector_db = FAISS.from_documents(chunks, embeddings)
    except Exception:
        try:
            # Tier 2: Try HuggingFace Embeddings
            from langchain_community.embeddings import HuggingFaceEmbeddings
            embeddings = HuggingFaceEmbeddings(model_name="all-MiniLM-L6-v2")
            vector_db = FAISS.from_documents(chunks, embeddings)
        except Exception:
            # Tier 3: High-Speed Zero-Dependency Embedding Fallback
            embeddings = LightweightFallbackEmbeddings(dim=384)
            vector_db = FAISS.from_documents(chunks, embeddings)

    # Sparse BM25 keyword index
    bm25_index = SimpleBM25(chunks)

    total_stats = {
        "files": file_stats,
        "total_documents": len(file_stats),
        "total_chunks": len(chunks),
        "total_words": sum(f["words"] for f in file_stats.values()),
    }

    return vector_db, bm25_index, chunks, len(chunks), total_stats


# Backward compatibility wrappers
def process_pdfs_to_vector_db(
    file_paths: Union[str, List[str]],
    chunk_size: int = 800,
    chunk_overlap: int = 100,
    embedding_model: str = "nomic-embed-text",
) -> Tuple[FAISS, int]:
    vector_db, _, _, chunk_count, _ = process_files_to_knowledge_base(
        file_paths=file_paths,
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        embedding_model=embedding_model,
    )
    return vector_db, chunk_count


def process_pdf_to_vector_db(file_path: str):
    db, _ = process_pdfs_to_vector_db(file_path)
    return db